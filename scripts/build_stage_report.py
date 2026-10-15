from __future__ import annotations

import shutil
import tempfile
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
COVER = ROOT / "Page de garde PFA-GINFO.docx"
OUT_DIR = ROOT / "output" / "docx"
OUT = OUT_DIR / "Rapport_PFA_Sentinel_PME.docx"


TITLE = "Système de détection et d’alerte précoce des cybermenaces ciblant les PME"
SHORT_TITLE = "Sentinel PME"
HOST = "CMRPI - Espace Maroc Cyberconfiance"
STUDENT = "Nom et prénom de l’étudiant(e) à compléter"
ACADEMIC_YEAR = "2026 / 2027"


def set_run_font(run, name: str = "Times New Roman", size: int | None = 12, bold: bool | None = None,
                 italic: bool | None = None, color: str | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_body(paragraph, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6) -> None:
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_after = Pt(space_after)
    fmt.space_before = Pt(0)
    for run in paragraph.runs:
        set_run_font(run)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False, fill: str | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    set_run_font(r, size=10, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        set_cell_shading(cell, fill)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "94A3B8")


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("Page ")
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_text, fld_end])


def add_toc(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-2" \h \z \u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Cliquez avec le bouton droit puis Mettre à jour le champ si nécessaire."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, placeholder, fld_end])


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    p.add_run(text)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_pr = p._p.get_or_add_pPr()
    outline = p_pr.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        p_pr.append(outline)
    outline.set(qn("w:val"), str(max(level - 1, 0)))
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 13, bold=True, color="1F2937")
    p.paragraph_format.space_before = Pt(12 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(6)


def add_para(doc: Document, text: str = "", align=WD_ALIGN_PARAGRAPH.JUSTIFY) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run_font(r)
    set_paragraph_body(p, align=align)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph()
        r = p.add_run(f"• {item}")
        set_run_font(r)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.paragraph_format.space_after = Pt(3)


def add_numbered(doc: Document, items: list[str]) -> None:
    for idx, item in enumerate(items, 1):
        p = doc.add_paragraph()
        r = p.add_run(f"{idx}. {item}")
        set_run_font(r)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.paragraph_format.space_after = Pt(3)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, fill="D9EAF7")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph()


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run_font(r, size=10, italic=True)
    p.paragraph_format.space_after = Pt(8)


def add_figure_box(doc: Document, text: str, caption: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_borders(table)
    set_row_cant_split(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F8FAFC")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_together = True
    r = p.add_run(text)
    set_run_font(r, name="Consolas", size=9)
    caption_p = cell.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_after = Pt(0)
    caption_p.paragraph_format.keep_together = True
    caption_run = caption_p.add_run(caption)
    set_run_font(caption_run, size=10, italic=True)
    doc.add_paragraph()


def add_inline_diagram(doc: Document, text: str, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_run_font(r, name="Consolas", size=9)
    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.keep_together = True
    caption_p.paragraph_format.space_after = Pt(6)
    cr = caption_p.add_run(caption)
    set_run_font(cr, size=10, italic=True)


def replace_cover_placeholders(doc: Document) -> None:
    replacements = {
        "Titre du projet": TITLE,
        "NOM & Prénom": STUDENT,
        "Nom de l’entreprise": HOST,
        "Année universitaire : 2026 / 2027": f"Année universitaire : {ACADEMIC_YEAR}",
    }
    for p in doc.paragraphs:
        original = p.text
        if original in replacements:
            for run in p.runs:
                run.text = ""
            r = p.add_run(replacements[original])
            set_run_font(r, size=18 if original == "Titre du projet" else 14, bold=original == "Titre du projet")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def patch_ooxml_text(path: Path, replacements: dict[str, str]) -> None:
    """Replace text that lives in OOXML shapes/tables not exposed by Document.paragraphs."""

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp) / path.name
        with ZipFile(path, "r") as zin, ZipFile(tmp_path, "w", compression=ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename.endswith(".xml"):
                    text = data.decode("utf-8")
                    for old, new in replacements.items():
                        text = text.replace(old, new)
                    data = text.encode("utf-8")
                zout.writestr(item, data)
        shutil.move(str(tmp_path), path)


def configure_content_section(doc: Document):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header_p = section.header.paragraphs[0]
    header_p.text = ""
    hr = header_p.add_run(f"{SHORT_TITLE} - Rapport de Projet de Fin d’Année")
    set_run_font(hr, size=9, italic=True, color="475569")
    header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p = section.footer.paragraphs[0]
    footer_p.text = ""
    add_page_number(footer_p)
    for run in footer_p.runs:
        set_run_font(run, size=9, color="475569")
    return section


def add_preliminaries(doc: Document) -> None:
    add_heading(doc, "Dédicace", 1)
    add_para(
        doc,
        "Je dédie ce travail à ma famille, à mes enseignants et à toutes les personnes qui m’ont accompagné dans mon parcours académique. Leur soutien, leurs conseils et leur confiance ont constitué un appui essentiel tout au long de ce projet.",
    )
    doc.add_page_break()

    add_heading(doc, "Remerciements", 1)
    add_para(
        doc,
        "Je tiens à exprimer ma reconnaissance à l’équipe du CMRPI - Espace Maroc Cyberconfiance pour le cadre de travail proposé et pour l’orientation donnée à ce sujet de cybersécurité appliquée aux petites et moyennes entreprises. Je remercie également mon encadrant pédagogique, les enseignants du département Informatique de l’École Nationale Supérieure des Mines de Rabat, ainsi que toutes les personnes ayant contribué par leurs remarques à l’amélioration de ce travail.",
    )
    add_para(
        doc,
        "Ce projet m’a permis de consolider mes compétences en développement full-stack, en Threat Intelligence, en conception d’API, en modélisation de données, en tests et en documentation technique.",
    )
    doc.add_page_break()

    add_heading(doc, "Fiche signalétique de l’organisme d’accueil", 1)
    add_table(
        doc,
        ["Élément", "Information"],
        [
            ["Organisme d’accueil", HOST],
            ["Domaine d’activité", "Recherche, innovation et accompagnement de projets numériques et cybersécurité"],
            ["Structure associée", "Espace Maroc Cyberconfiance"],
            ["Sujet du stage", TITLE],
            ["Type de stage", "Projet de Fin d’Année"],
            ["Département", "Informatique"],
            ["Filière", "Génie Informatique"],
            ["Année universitaire", ACADEMIC_YEAR],
            ["Informations administratives", "À compléter selon les données officielles de l’organisme"],
        ],
        widths=[2.2, 4.3],
    )
    doc.add_page_break()

    add_heading(doc, "Résumé", 1)
    add_para(
        doc,
        "Les petites et moyennes entreprises sont fortement exposées aux cybermenaces, alors qu’elles disposent rarement d’un centre opérationnel de sécurité, d’un SIEM complet ou d’une équipe spécialisée. Le stage présenté dans ce rapport porte sur la conception et la réalisation de Sentinel PME, une plateforme web de veille, de détection et d’alerte précoce des cybermenaces ciblant les PME. La solution collecte des indicateurs depuis des sources de Threat Intelligence, les normalise, calcule un score de risque explicable, génère des alertes et propose des fiches réflexes adaptées aux capacités d’une PME. Le projet repose sur une architecture full-stack composée d’un backend FastAPI, d’une base SQL, de collecteurs Python, d’un moteur de scoring et d’un dashboard React. Des extensions de niveau entreprise ont été ajoutées, notamment un workflow SOC léger, un export SIEM JSON, un inventaire des actifs, un rapprochement entre logiciels et CVE, ainsi qu’un enrichissement par EPSS et CISA KEV. La validation backend a été assurée par une suite de tests automatisés, tandis que le frontend a été vérifié par compilation TypeScript et build Vite. Le résultat obtenu constitue un MVP crédible, extensible et aligné avec les besoins d’une PME en matière de priorisation et de réponse rapide.",
    )
    add_para(doc, "Mots clés : cybersécurité, PME, Threat Intelligence, FastAPI, React, SIEM-light, SOC, CVE.")

    add_heading(doc, "Abstract", 1)
    add_para(
        doc,
        "Small and medium-sized enterprises are increasingly exposed to cyber threats while often lacking a dedicated security operations center, a full SIEM platform or specialized staff. This internship report presents the design and implementation of Sentinel PME, a web platform for cyber threat intelligence, early warning and risk prioritization for SMEs. The solution collects threat indicators from public intelligence sources, normalizes heterogeneous data, computes an explainable risk score, generates alerts and provides practical response sheets. The project follows a full-stack architecture combining a FastAPI backend, a SQL database, Python collectors, a rule-based scoring engine and a React dashboard. Enterprise-oriented extensions were introduced, including a lightweight SOC workflow, SIEM-friendly JSON export, asset inventory, software-to-CVE matching and enrichment through EPSS and CISA KEV. Backend validation was performed through automated tests, and the frontend was verified through TypeScript and Vite production build. The resulting MVP is realistic, extensible and aligned with the operational needs of SMEs.",
    )
    add_para(doc, "Keywords: cybersecurity, SME, Threat Intelligence, FastAPI, React, SIEM-light, SOC, CVE.")

    add_heading(doc, "ملخص", 1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(
        "تعالج هذه الدراسة موضوع بناء منصة خفيفة لمراقبة التهديدات السيبرانية الموجهة للمقاولات الصغيرة والمتوسطة. تعتمد المنصة على جمع مؤشرات التهديد من مصادر خارجية، ثم توحيدها وتحليلها وحساب درجة خطورة قابلة للتفسير. كما توفر لوحة قيادة لعرض النتائج، وتنبيهات للحالات الحرجة، وبطاقات إرشادية تساعد المؤسسة على اتخاذ إجراءات عملية. تم تطوير المشروع باستعمال FastAPI في الواجهة الخلفية وReact في الواجهة الأمامية، مع قاعدة بيانات SQL ومحركات لجمع وتحليل البيانات. كما تم تعزيز الحل بوظائف قريبة من بيئة SOC وSIEM مثل تتبع حالة التنبيهات، تصدير الأحداث بصيغة JSON، وجرد الأصول وربطها بالثغرات."
    )
    set_run_font(r, name="Arial", size=12)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    add_para(doc, "الكلمات المفتاحية : الأمن السيبراني، المقاولات الصغيرة والمتوسطة، الذكاء التهديدي، التنبيهات، الثغرات.")
    doc.add_page_break()

    add_heading(doc, "Table des matières", 1)
    add_toc(doc.add_paragraph())
    doc.add_page_break()

    add_heading(doc, "Liste des acronymes et abréviations", 1)
    add_table(
        doc,
        ["Acronyme", "Signification"],
        [
            ["API", "Application Programming Interface"],
            ["CISA KEV", "Cybersecurity and Infrastructure Security Agency Known Exploited Vulnerabilities"],
            ["CPE", "Common Platform Enumeration"],
            ["CVE", "Common Vulnerabilities and Exposures"],
            ["CVSS", "Common Vulnerability Scoring System"],
            ["EDR", "Endpoint Detection and Response"],
            ["EPSS", "Exploit Prediction Scoring System"],
            ["IOC", "Indicator of Compromise"],
            ["JWT", "JSON Web Token"],
            ["MVP", "Minimum Viable Product"],
            ["NVD", "National Vulnerability Database"],
            ["PME", "Petite et Moyenne Entreprise"],
            ["RBAC", "Role-Based Access Control"],
            ["SIEM", "Security Information and Event Management"],
            ["SOC", "Security Operations Center"],
        ],
        widths=[1.5, 5.0],
    )

    add_heading(doc, "Liste des figures", 1)
    add_table(
        doc,
        ["N°", "Titre"],
        [
            ["Figure 1", "Architecture générale de Sentinel PME"],
            ["Figure 2", "Flux SOC/SIEM proposé"],
        ],
    )

    add_heading(doc, "Liste des tableaux", 1)
    add_table(
        doc,
        ["N°", "Titre"],
        [
            ["Tableau 1", "Sources de Threat Intelligence exploitées"],
            ["Tableau 2", "Niveaux de risque utilisés par le moteur de scoring"],
            ["Tableau 3", "Stack technologique du projet"],
            ["Tableau 4", "Synthèse des tests et validations"],
        ],
    )
    doc.add_page_break()


def add_introduction(doc: Document) -> None:
    add_heading(doc, "Introduction générale", 1)
    add_para(
        doc,
        "La transformation numérique a rendu les petites et moyennes entreprises dépendantes de services connectés : messagerie professionnelle, applications cloud, sites web, outils de gestion, postes de travail et serveurs exposés. Cette dépendance augmente leur surface d’attaque, alors que leurs moyens humains et financiers restent souvent limités. Beaucoup de PME ne disposent pas d’un SOC, d’un SIEM complet ou d’une équipe cybersécurité permanente. Elles ont donc besoin d’outils simples, compréhensibles et orientés action.",
    )
    add_para(
        doc,
        "Le sujet de ce Projet de Fin d’Année porte sur la conception et le développement d’un système de détection et d’alerte précoce des cybermenaces ciblant les PME. Le projet, nommé Sentinel PME, vise à transformer des données de Threat Intelligence en informations exploitables : indicateurs normalisés, niveau de risque, alertes, recommandations et fiches réflexes.",
    )
    add_para(
        doc,
        "La problématique traitée peut être formulée ainsi : comment fournir à une PME une plateforme légère capable de collecter, prioriser et présenter les menaces cyber pertinentes, sans imposer la complexité d’un SIEM d’entreprise ?",
    )
    add_para(
        doc,
        "La méthodologie suivie s’appuie sur une démarche incrémentale : analyse du besoin, état de l’art, cadrage du MVP, conception de l’architecture, développement du backend, réalisation du dashboard, intégration des collecteurs, ajout du scoring, validation par tests et documentation. Le rapport est organisé en cinq chapitres : contexte et organisme d’accueil, analyse du besoin, conception, réalisation technique, puis tests et discussion.",
    )


def add_chapter_1(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "Chapitre 1 : Organisme d’accueil et contexte du projet", 1)
    add_para(doc, "Objectif du chapitre : présenter l’environnement du stage, le contexte de cybersécurité des PME et la finalité du projet.")
    add_heading(doc, "1.1 Présentation de l’organisme d’accueil", 2)
    add_para(
        doc,
        "Le stage s’inscrit dans le cadre du CMRPI - Espace Maroc Cyberconfiance. Le sujet proposé se situe à l’intersection de la cybersécurité, de l’intelligence artificielle appliquée, de la Threat Intelligence et du développement web. L’objectif pédagogique est de produire un prototype utile, démontrable et extensible, tout en respectant les contraintes réalistes d’un Projet de Fin d’Année.",
    )
    add_para(
        doc,
        "L’organisme d’accueil joue ici un rôle d’encadrement, de définition du besoin et d’orientation technique. Le projet répond à un axe stratégique lié à la cybersécurité des PME et à la valorisation de solutions accessibles pour des organisations qui ne disposent pas d’une infrastructure de sécurité lourde.",
    )
    add_heading(doc, "1.2 Contexte de cybersécurité des PME", 2)
    add_para(
        doc,
        "Les PME sont exposées à des menaces comparables à celles des grandes organisations : phishing, ransomwares, malwares, exploitation de vulnérabilités, attaques par force brute ou réutilisation d’identifiants compromis. Leur difficulté principale n’est pas seulement l’existence de ces menaces, mais l’absence d’un mécanisme simple de priorisation et de réponse.",
    )
    add_para(
        doc,
        "Dans une grande entreprise, les données de sécurité sont généralement traitées par des outils spécialisés comme les SIEM, les EDR, les scanners de vulnérabilités et les plateformes SOAR. Dans une PME, ces solutions peuvent être trop coûteuses, trop complexes ou mal adaptées au niveau de maturité disponible. Sentinel PME se positionne donc comme une brique intermédiaire : plus structurée qu’un simple script, mais moins lourde qu’un SIEM complet.",
    )
    add_heading(doc, "1.3 Objectifs du stage", 2)
    add_bullets(
        doc,
        [
            "Développer une plateforme web capable de collecter des données de Threat Intelligence.",
            "Normaliser des indicateurs hétérogènes tels que les IP, domaines, URLs, CVE et hashes.",
            "Calculer un score de risque explicable compris entre 0 et 100.",
            "Générer des alertes pour les menaces élevées et critiques.",
            "Afficher les résultats dans un dashboard compréhensible pour une PME.",
            "Fournir des fiches réflexes et une base d’intégration SOC/SIEM.",
        ],
    )


def add_chapter_2(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "Chapitre 2 : Analyse du besoin et état de l’art", 1)
    add_para(doc, "Objectif du chapitre : analyser les cybermenaces visées, les sources de données utilisées et le positionnement du projet par rapport aux solutions existantes.")
    add_heading(doc, "2.1 Problématique", 2)
    add_para(
        doc,
        "Une PME doit pouvoir répondre rapidement à plusieurs questions : quelle menace est prioritaire, quel indicateur doit être bloqué, quelle vulnérabilité doit être corrigée en premier et quelle action doit être menée en cas d’alerte. Sans outil de centralisation, les flux de Threat Intelligence restent difficiles à exploiter, car chaque source possède son propre format, son vocabulaire et son niveau de fiabilité.",
    )
    add_heading(doc, "2.2 Cybermenaces ciblées", 2)
    add_table(
        doc,
        ["Menace", "Description", "Impact pour une PME"],
        [
            ["Phishing", "Emails ou sites frauduleux destinés à voler des identifiants.", "Vol de comptes, fraude et compromission initiale."],
            ["Ransomware", "Malware chiffrant les données puis réclamant une rançon.", "Arrêt d’activité, perte de données et atteinte à l’image."],
            ["Malware", "Logiciel malveillant servant à espionner, voler ou détruire.", "Exfiltration, sabotage ou persistance sur le réseau."],
            ["Brute force", "Tentatives répétées de connexion.", "Compromission de comptes exposés."],
            ["Credential stuffing", "Réutilisation d’identifiants déjà divulgués.", "Prise de contrôle de comptes professionnels."],
            ["Vulnérabilité critique", "Faille CVE exploitable sur un système non corrigé.", "Intrusion, exécution de code ou fuite de données."],
        ],
        widths=[1.5, 2.7, 2.5],
    )
    add_heading(doc, "2.3 Sources de Threat Intelligence", 2)
    add_table(
        doc,
        ["Source", "Type de données", "Utilisation dans Sentinel PME"],
        [
            ["NVD / CVE", "Vulnérabilités, CVSS, configurations CPE", "Identifier et prioriser les vulnérabilités critiques."],
            ["URLhaus", "URLs de distribution de malware", "Détecter les liens malveillants actifs ou récents."],
            ["AbuseIPDB", "Adresses IP signalées", "Identifier des IP suspectes ou hostiles."],
            ["AlienVault OTX", "Indicateurs communautaires", "Collecter domaines, IP, URLs, hashes et CVE."],
            ["EPSS", "Probabilité d’exploitation", "Enrichir la priorisation des CVE."],
            ["CISA KEV", "Vulnérabilités exploitées dans la nature", "Renforcer l’urgence de traitement des CVE confirmées exploitées."],
        ],
        widths=[1.4, 2.2, 3.0],
    )
    add_heading(doc, "2.4 Positionnement par rapport aux SIEM et SOC", 2)
    add_para(
        doc,
        "Sentinel PME ne remplace pas un SIEM. Un SIEM collecte et corrèle massivement les journaux internes d’une organisation. Un SOC combine des analystes, des procédures et des outils pour surveiller, qualifier et répondre aux incidents. Sentinel PME agit plutôt comme une brique de Threat Intelligence et d’alerte précoce : elle produit des alertes enrichies pouvant être exploitées par un mini-SOC ou exportées vers un SIEM externe.",
    )


def add_chapter_3(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "Chapitre 3 : Conception de la solution", 1)
    add_para(doc, "Objectif du chapitre : présenter l’architecture, le modèle de données, le moteur de scoring et les choix de conception retenus.")
    add_heading(doc, "3.1 Architecture générale", 2)
    add_figure_box(
        doc,
        "Threat Intelligence APIs\n        ↓\nPython Collectors\n        ↓\nNormalizer\n        ↓\nRisk Scoring Engine\n        ↓\nDatabase\n        ↓\nFastAPI Backend\n        ↓\nReact Dashboard\n        ↓\nAlerts + Fiches Réflexes",
        "Figure 1 : Architecture générale de Sentinel PME",
    )
    add_para(
        doc,
        "L’architecture est organisée en couches. Les collecteurs récupèrent les données externes et les convertissent vers un format commun. Le moteur de scoring attribue un niveau de risque. Les services métier assurent la déduplication, la création d’alertes et l’association des recommandations. L’API FastAPI expose les données au dashboard React.",
    )
    add_heading(doc, "3.2 Modèle de données principal", 2)
    add_para(
        doc,
        "La donnée centrale du projet est la menace normalisée. Elle contient un indicateur, son type, la source, la sévérité, le score de risque, le niveau de risque, une description, une action recommandée et des dates d’observation. Les alertes référencent les menaces dont le score dépasse le seuil défini. Les actifs et logiciels permettent de relier certaines CVE à un parc informatique déclaré.",
    )
    add_table(
        doc,
        ["Entité", "Rôle", "Champs principaux"],
        [
            ["Threat", "Menace normalisée et scorée", "indicator, indicator_type, threat_type, source, severity, risk_score, risk_level, score_factors"],
            ["Alert", "Alerte générée à partir d’une menace", "level, status, assigned_to, resolution_notes, acknowledged"],
            ["Asset", "Équipement ou application de la PME", "hostname, type, environnement, criticité, exposition Internet"],
            ["Software", "Logiciel installé sur un actif", "name, vendor, version, CPE, PURL"],
            ["AuditLog", "Traçabilité des actions sensibles", "action, actor, target, timestamp"],
        ],
        widths=[1.3, 2.1, 3.3],
    )
    add_heading(doc, "3.3 Moteur de scoring", 2)
    add_para(
        doc,
        "Le scoring est déterministe et explicable. Il ne cherche pas à prédire scientifiquement la probabilité d’un incident ; il sert à ordonner les menaces selon leur urgence opérationnelle. Le score est calculé à partir de facteurs additifs et plafonné à 100.",
    )
    add_table(
        doc,
        ["Score", "Niveau", "Traitement attendu"],
        [
            ["0 - 39", "Faible", "Surveillance et conservation."],
            ["40 - 69", "Moyen", "Vérification dans le cycle normal."],
            ["70 - 89", "Élevé", "Alerte et traitement rapide."],
            ["90 - 100", "Critique", "Alerte critique et action immédiate."],
        ],
        widths=[1.2, 1.4, 4.0],
    )
    add_heading(doc, "3.4 Extension SOC/SIEM", 2)
    add_figure_box(
        doc,
        "Alerte nouvelle -> Investigation -> Résolution ou faux positif\n              ↓\n       Export JSON SIEM / Webhook SOC",
        "Figure 2 : Flux SOC/SIEM proposé",
    )
    add_para(
        doc,
        "Pour rapprocher le projet d’un contexte entreprise, les alertes ont été enrichies par un mini-workflow SOC. Un analyste peut suivre le statut, affecter une alerte et documenter la résolution. Un endpoint exporte l’événement au format JSON afin qu’il puisse être consommé par Wazuh, Elastic Security, Splunk ou un outil SOAR.",
    )


def add_chapter_4(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "Chapitre 4 : Réalisation technique", 1)
    add_para(doc, "Objectif du chapitre : décrire les modules développés, les technologies utilisées et les principaux choix d’implémentation.")
    add_heading(doc, "4.1 Stack technologique", 2)
    add_table(
        doc,
        ["Couche", "Technologies", "Rôle"],
        [
            ["Backend", "Python, FastAPI, Pydantic, SQLAlchemy", "API REST, validation et logique métier."],
            ["Base de données", "SQLite en développement, PostgreSQL cible", "Persistance des menaces, alertes, utilisateurs et actifs."],
            ["Collecte", "httpx, collecteurs Python, APScheduler", "Interrogation des sources et collecte planifiée."],
            ["Frontend", "React, TypeScript, Vite, Tailwind CSS, Axios", "Dashboard, filtres, détails et navigation."],
            ["Sécurité", "JWT, bcrypt, RBAC, slowapi", "Authentification, rôles et limitation de débit."],
            ["Qualité", "pytest, Vitest, TypeScript, Docker, GitHub Actions", "Tests, build, CI et industrialisation."],
        ],
        widths=[1.4, 2.4, 2.8],
    )
    add_heading(doc, "4.2 Backend FastAPI", 2)
    add_para(
        doc,
        "Le backend expose des routes pour l’authentification, les menaces, les alertes, les collecteurs, les fiches réflexes, les actifs et le journal d’audit. Les routes HTTP délèguent la logique aux services internes afin de séparer les responsabilités. Les schémas Pydantic valident les entrées et structurent les sorties.",
    )
    add_bullets(
        doc,
        [
            "Création automatique d’un compte administrateur au premier démarrage si aucun utilisateur n’existe.",
            "Contrôle d’accès par rôles : administrateur, analyste et responsable PME.",
            "Rate limiting sur les endpoints d’écriture.",
            "En-têtes de sécurité HTTP et CORS configuré.",
            "Migrations Alembic pour suivre les évolutions du schéma.",
        ],
    )
    add_heading(doc, "4.3 Collecteurs et normalisation", 2)
    add_para(
        doc,
        "Les collecteurs transforment les réponses des sources de Threat Intelligence vers un modèle `ThreatInput`. Cette normalisation évite que le reste de l’application dépende des formats propres à NVD, URLhaus, AbuseIPDB ou OTX. Le service de menace réalise ensuite la déduplication par paire indicateur/source.",
    )
    add_heading(doc, "4.4 Dashboard React", 2)
    add_para(
        doc,
        "Le frontend fournit une interface protégée par authentification. Il comprend une vue d’ensemble, un registre des menaces, une page de détail avec explication du score, un centre d’alertes, une vue du parc informatique, les fiches réflexes, les sources de collecte et une page de présentation. Le client Axios centralise les appels API et la gestion du jeton JWT.",
    )
    add_heading(doc, "4.5 Inventaire des actifs et matching CVE", 2)
    add_para(
        doc,
        "L’inventaire des actifs permet de relier une menace de type CVE à un logiciel réellement présent dans le parc déclaré par la PME. Les actifs possèdent un hostname, une criticité, un environnement et une exposition Internet. Les logiciels peuvent contenir un CPE 2.3, ce qui améliore la précision du rapprochement avec les configurations NVD.",
    )
    add_para(
        doc,
        "Le moteur de matching adopte une stratégie conservatrice. Il distingue les correspondances confirmées, potentielles et non affectées. Une correspondance n’est confirmée que si le produit, le CPE et la version permettent une conclusion suffisamment fiable.",
    )


def add_chapter_5(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "Chapitre 5 : Tests, validation et discussion", 1)
    add_para(doc, "Objectif du chapitre : présenter les contrôles réalisés, les résultats obtenus, les limites du MVP et les perspectives.")
    add_heading(doc, "5.1 Stratégie de validation", 2)
    add_para(
        doc,
        "La validation du projet repose sur des tests automatisés backend, une compilation du frontend, des scénarios fonctionnels et une vérification de la documentation. Les tests backend utilisent une base SQLite dédiée et ne dépendent pas d’appels externes. Cela permet de valider la logique métier de manière reproductible.",
    )
    add_table(
        doc,
        ["Contrôle", "Résultat observé", "Commentaire"],
        [
            ["Tests backend pytest", "44 tests passés", "Validation API, scoring, collecteurs, actifs, enrichissement et seed de démonstration."],
            ["Build frontend", "Succès", "Compilation TypeScript et génération du bundle Vite."],
            ["Tests frontend Vitest", "À relancer en environnement stable", "La commande lancée dans l’environnement courant n’a pas terminé avant le délai disponible."],
            ["Documentation", "Présente dans `docs/`", "Cahier des charges, architecture, scoring, API, SOC/SIEM, CI/CD et enrichissement CVE."],
        ],
        widths=[1.7, 1.8, 3.2],
    )
    add_heading(doc, "5.2 Scénarios fonctionnels validés", 2)
    add_numbered(
        doc,
        [
            "Ajout d’une menace critique puis génération automatique d’une alerte.",
            "Analyse d’un indicateur et calcul d’un score de risque explicable.",
            "Liste filtrable des menaces avec pagination et recherche.",
            "Consultation d’une fiche réflexe selon le type de menace.",
            "Mise à jour d’une alerte dans un workflow SOC.",
            "Export d’une alerte au format SIEM-friendly.",
            "Création et import d’actifs avec validation des champs.",
            "Rapprochement entre logiciels déclarés et vulnérabilités CVE stockées.",
        ],
    )
    add_heading(doc, "5.3 Analyse des résultats", 2)
    add_para(
        doc,
        "Le projet atteint l’objectif principal : produire une plateforme fonctionnelle de veille et d’alerte précoce pour PME. Le passage d’un simple prototype à une architecture backend/frontend renforce la crédibilité du projet. L’ajout du workflow SOC, de l’export SIEM et de l’inventaire des actifs rapproche la solution d’un contexte entreprise tout en conservant un périmètre maîtrisé.",
    )
    add_para(
        doc,
        "Le moteur de scoring présente l’avantage d’être transparent : chaque score peut être expliqué par des facteurs tels que la sévérité, le type d’indicateur, la source, la récence, EPSS, CISA KEV et la corrélation multi-sources. Cette explicabilité est essentielle pour une PME qui doit comprendre pourquoi une menace devient prioritaire.",
    )
    add_heading(doc, "5.4 Limites", 2)
    add_bullets(
        doc,
        [
            "La solution ne remplace pas un SIEM complet et ne collecte pas massivement les logs internes.",
            "Le scoring reste fondé sur des règles ; il doit être calibré avec des retours analystes réels.",
            "La qualité dépend de la disponibilité et de la fiabilité des sources publiques.",
            "L’inventaire des actifs nécessite une donnée déclarative propre pour produire un matching CVE pertinent.",
            "Les canaux email, Telegram ou Slack restent des extensions à finaliser selon l’environnement cible.",
        ],
    )
    add_heading(doc, "5.5 Perspectives", 2)
    add_bullets(
        doc,
        [
            "Ajouter un modèle Scikit-learn pour compléter la classification et comparer ses résultats au scoring à règles.",
            "Développer un connecteur Wazuh ou Elastic Security dédié.",
            "Mettre en place des notifications multicanal avec préférences par criticité.",
            "Générer automatiquement des rapports PDF de situation cyber pour le responsable PME.",
            "Améliorer la corrélation multi-sources et la prise en compte du contexte métier des actifs.",
            "Déployer une version PostgreSQL et Docker Compose prête pour un environnement de démonstration entreprise.",
        ],
    )


def add_conclusion_and_refs(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "Conclusion générale", 1)
    add_para(
        doc,
        "Ce Projet de Fin d’Année a permis de concevoir et développer Sentinel PME, une plateforme de Threat Intelligence et d’alerte précoce destinée aux petites et moyennes entreprises. La solution répond à une problématique concrète : aider une PME à comprendre les menaces pertinentes, à les prioriser et à agir rapidement sans disposer d’un SOC ou d’un SIEM complet.",
    )
    add_para(
        doc,
        "Le travail réalisé couvre l’analyse du besoin, l’état de l’art, la conception d’une architecture modulaire, le développement d’un backend FastAPI, la réalisation d’un dashboard React, l’intégration de collecteurs, le scoring explicable, les alertes, les fiches réflexes, l’inventaire des actifs, l’enrichissement des CVE et une couche SOC/SIEM-light. Les tests backend et le build frontend confirment la stabilité du socle développé.",
    )
    add_para(
        doc,
        "Les principales limites concernent l’absence de collecte massive de logs internes, le caractère encore perfectible du scoring, la dépendance aux sources publiques et le besoin de calibrage avec des données réelles d’entreprise. Ces limites sont cohérentes avec le périmètre d’un MVP PFA et ouvrent des pistes pertinentes pour une évolution vers un PFE ou une solution entreprise.",
    )
    add_para(
        doc,
        "À terme, Sentinel PME peut évoluer vers une plateforme plus complète intégrant des connecteurs SIEM dédiés, un modèle Machine Learning explicable, des notifications multicanal, des rapports automatisés et une meilleure corrélation entre les menaces, les actifs et le contexte métier.",
    )

    doc.add_page_break()
    add_heading(doc, "Références", 1)
    refs = [
        "ENSMR, Guide de rédaction des rapports, Département Informatique, Année universitaire 2026-2027.",
        "National Vulnerability Database, NVD Developers - Vulnerabilities API, https://nvd.nist.gov/developers/vulnerabilities, consulté le 14/10/2026.",
        "abuse.ch, URLhaus Community API, https://urlhaus.abuse.ch/api/, consulté le 14/10/2026.",
        "FIRST, EPSS API documentation, https://api.first.org/epss/, consulté le 14/10/2026.",
        "CISA, Known Exploited Vulnerabilities Catalog, https://www.cisa.gov/known-exploited-vulnerabilities-catalog, consulté le 14/10/2026.",
        "FastAPI Documentation, https://fastapi.tiangolo.com/.",
        "React Documentation, https://react.dev/.",
        "SQLAlchemy Documentation, https://docs.sqlalchemy.org/.",
        "Documentation interne du projet Sentinel PME : README, cahier des charges, architecture, scoring, API, tests, SOC/SIEM et inventaire des actifs.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        r = p.add_run(f"[{refs.index(ref) + 1}] {ref}")
        set_run_font(r)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    doc.add_page_break()
    add_heading(doc, "Annexes", 1)
    add_heading(doc, "Annexe A : Commandes de lancement", 2)
    add_para(doc, "Backend :")
    add_figure_box(
        doc,
        "cd backend\n.\\.venv\\Scripts\\python.exe -m uvicorn app.main:app --reload",
        "Extrait A.1 : Commande de lancement du backend",
    )
    add_para(doc, "Frontend :")
    add_figure_box(
        doc,
        "cd frontend\nnpm run dev",
        "Extrait A.2 : Commande de lancement du frontend",
    )
    add_heading(doc, "Annexe B : Commandes de validation", 2)
    add_figure_box(
        doc,
        "cd backend\n.\\.venv\\Scripts\\python.exe -m pytest -q\n\ncd frontend\nnpm run build",
        "Extrait B.1 : Commandes de test et de build",
    )


def setup_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(COVER, OUT)
    doc = Document(OUT)
    replace_cover_placeholders(doc)
    setup_styles(doc)
    configure_content_section(doc)
    add_preliminaries(doc)
    add_introduction(doc)
    add_chapter_1(doc)
    add_chapter_2(doc)
    add_chapter_3(doc)
    add_chapter_4(doc)
    add_chapter_5(doc)
    add_conclusion_and_refs(doc)

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    doc.save(OUT)
    patch_ooxml_text(
        OUT,
        {
            "Titre du projet": TITLE,
            "Nom de l’entreprise": HOST,
            "NOM &amp; Prénom": STUDENT,
            "NOM & Prénom": STUDENT,
        },
    )
    print(OUT)


if __name__ == "__main__":
    main()
